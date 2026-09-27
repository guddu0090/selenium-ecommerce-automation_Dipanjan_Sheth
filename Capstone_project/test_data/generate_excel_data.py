"""
Run once to create test_data.xlsx alongside test_data.json.
This demonstrates reading test data from Excel as an alternative to JSON.
"""
import json
import os
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))

with open(os.path.join(HERE, "test_data.json")) as f:
    data = json.load(f)

wb = Workbook()

ws1 = wb.active
ws1.title = "General"
ws1.append(["Key", "Value"])
ws1.append(["base_url", data["base_url"]])
ws1.append(["search_product", data["search_product"]])
ws1.append(["quantity_to_set", data["quantity_to_set"]])

ws2 = wb.create_sheet("User")
ws2.append(["Field", "Value"])
for k, v in data["user"].items():
    ws2.append([k, v])

ws3 = wb.create_sheet("SignupDetails")
ws3.append(["Field", "Value"])
for k, v in data["signup_details"].items():
    ws3.append([k, v])

out_path = os.path.join(HERE, "test_data.xlsx")
wb.save(out_path)
print(f"Excel test data written to {out_path}")
